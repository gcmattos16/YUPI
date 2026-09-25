import os
from datetime import datetime


def exportar_excel(caminho, resultados, metadados=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Conferencia"

    ws.append(["YUPI - Conferência de Consultas"])
    ws["A1"].font = Font(size=16, bold=True)
    ws.merge_cells("A1:F1")

    if metadados:
        ws.append(["Gerado em", datetime.now().strftime("%d/%m/%Y %H:%M")])
        ws.append(["SIRESP", metadados.get("siresp", "")])
        ws.append(["Analítico", metadados.get("analitico", "")])
    ws.append([])
    ws.append(["Especialidade", "SIRESP", "Analítico", "Diferença", "Status", "Composição"])

    for item in resultados:
        detalhes = []
        for origem, chave in (("SIRESP", "detalhes_siresp"), ("Analítico", "detalhes_analitico")):
            partes = [f"{d.get('especialidade')}: {d.get('quantidade', 0)}" for d in item.get(chave, [])]
            if partes:
                detalhes.append(f"{origem}: " + " + ".join(partes))
        ws.append([
            item.get("especialidade", ""),
            item.get("siresp", 0),
            item.get("analitico", 0),
            item.get("diferenca", 0),
            item.get("status", ""),
            " | ".join(detalhes),
        ])

    for cell in ws[ws.max_row - len(resultados)]:
        cell.font = Font(bold=True)
    widths = {"A": 38, "B": 12, "C": 12, "D": 12, "E": 18, "F": 90}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(caminho)
    return caminho


def exportar_pdf(caminho, resultados, metadados=None):
    import fitz

    doc = fitz.open()
    pagina = None
    y = 0

    def nova_pagina():
        nonlocal pagina, y
        pagina = doc.new_page(width=842, height=595)  # A4 paisagem
        y = 38
        pagina.insert_text((36, y), "YUPI - Conferência de Consultas", fontsize=16)
        y += 24
        if metadados:
            pagina.insert_text((36, y), f"SIRESP: {metadados.get('siresp', '')}", fontsize=8)
            y += 12
            pagina.insert_text((36, y), f"Analítico: {metadados.get('analitico', '')}", fontsize=8)
            y += 18
        pagina.insert_text((36, y), "Especialidade", fontsize=9)
        pagina.insert_text((390, y), "SIRESP", fontsize=9)
        pagina.insert_text((455, y), "Analítico", fontsize=9)
        pagina.insert_text((525, y), "Dif.", fontsize=9)
        pagina.insert_text((580, y), "Status", fontsize=9)
        y += 16

    nova_pagina()
    for item in resultados:
        if y > 555:
            nova_pagina()
        nome = str(item.get("especialidade", ""))[:52]
        pagina.insert_text((36, y), nome, fontsize=8)
        pagina.insert_text((400, y), str(item.get("siresp", 0)), fontsize=8)
        pagina.insert_text((465, y), str(item.get("analitico", 0)), fontsize=8)
        pagina.insert_text((530, y), str(item.get("diferenca", 0)), fontsize=8)
        status = "OK" if item.get("diferenca", 0) == 0 else "DIVERGENCIA"
        pagina.insert_text((580, y), status, fontsize=8)
        y += 14

    doc.save(caminho)
    doc.close()
    return caminho
