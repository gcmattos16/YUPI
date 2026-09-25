from __future__ import annotations
from collections import defaultdict
from datetime import datetime, date, timedelta
import re, unicodedata
from openpyxl import load_workbook

# Abas auxiliares que não representam exame.
IGNORE_SHEETS = {'PLANILHA WEB SADT','PLAN1','GRAF1','GRAF 1','MENU','INICIO','INÍCIO'}

# Nome da aba/título -> nome canônico usado pelo YUPI.
ALIASES = {
    'NASO':'NASOFIBROSCOPIA', 'NASOFIBROSCOPIA':'NASOFIBROSCOPIA',
    'RAIO X':'RX', 'RAIO-X':'RX', 'RX':'RX',
    'BIOPSIA PROSTATA':'BIOPSIA DE PROSTATA', 'BIOPSIA DE PROSTATA':'BIOPSIA DE PROSTATA',
    'PFP':'PROVA DE FUNCAO PULMONAR', 'PROVA DE FUNCAO':'PROVA DE FUNCAO PULMONAR',
    'PROVA DE FUNCAO PULMONAR':'PROVA DE FUNCAO PULMONAR',
    'COLONO':'COLONOSCOPIA', 'COLONOSCOPIA':'COLONOSCOPIA',
    'RETO':'RETOSSIGMOIDOSCOPIA', 'RETOSSIGMOIDOSCOPIA':'RETOSSIGMOIDOSCOPIA',
    'ECOCARDIO':'ECOCARDIOGRAFIA', 'ECOCARDIOGRAFIA':'ECOCARDIOGRAFIA',
    'ECG':'ELETROCARDIOGRAMA', 'ELETROCARDIOGRAMA':'ELETROCARDIOGRAMA',
    'ELETROENCEFALO':'ELETROENCEFALOGRAMA', 'ELETROENCEFALOGRAMA':'ELETROENCEFALOGRAMA',
    'ELETRONEURO':'ELETRONEUROMIOGRAFIA', 'ELETRONEUROMIOGRAFIA':'ELETRONEUROMIOGRAFIA',
    'USG OLHO':'ULTRASSONOGRAFIA DE OLHO',
    'USG CONVENCIONAL':'USG CONVENCIONAL', 'USG DOPPLER':'USG DOPPLER',
    'OCT':'OCT', 'RETINOGRAFIA':'RETINOGRAFIA', 'BIOMETRIA':'BIOMETRIA',
    'CAMPIMETRIA':'CAMPIMETRIA', 'DENSITOMETRIA':'DENSITOMETRIA', 'MAMOGRAFIA':'MAMOGRAFIA',
    'AUDIOMETRIA':'AUDIOMETRIA', 'IMITANCIOMETRIA':'IMITANCIOMETRIA', 'LOGOAUDIOMETRIA':'LOGOAUDIOMETRIA',
    'BERA':'BERA', 'EMISSOES':'EMISSOES OTOACUSTICAS', 'OTONEURO':'OTONEUROLOGICO',
    'COLPOSCOPIA':'COLPOSCOPIA', 'HOLTER':'HOLTER', 'MAPA':'MAPA', 'URODINAMICA':'URODINAMICA',
    'UROFLUXOMETRIA':'UROFLUXOMETRIA', 'TESTE ERGOMETRICO':'TESTE ERGOMETRICO',
    'GONIOSCOPIA':'GONIOSCOPIA', 'PAM':'PAM', 'CARTOES DE TELLER':'CARTOES DE TELLER',
    'PAQUIMETRIA':'PAQUIMETRIA', 'TESTE ORTOPTICO':'TESTE ORTOPTICO', 'TOPOGRAFIA':'TOPOGRAFIA',
    'BIOMICROSCOPIA':'BIOMICROSCOPIA', 'MAPEAMENTO DE RETINA':'MAPEAMENTO DE RETINA',
    'FUNDOSCOPIA':'FUNDOSCOPIA', 'TONOMETRIA':'TONOMETRIA', 'ENDOSCOPIA':'ENDOSCOPIA',
    'ANUSCOPIA':'ANUSCOPIA', 'TESTE DE CONTATO':'TESTE DE CONTATO',
}

def _norm(s):
    s=unicodedata.normalize('NFKD', str(s or '')).encode('ascii','ignore').decode().upper().strip()
    return re.sub(r'\s+',' ',s)

def _canonical(s):
    n=_norm(s)
    return ALIASES.get(n,n)

def _exam_name(sheet, ws):
    # O nome da aba é a fonte mais estável; o título é usado para confirmar/melhorar.
    sheet_key=_canonical(sheet)
    title=''
    maxr=ws.max_row or 0
    for row in ws.iter_rows(min_row=1,max_row=min(maxr,6),values_only=True):
        for v in row:
            if isinstance(v,str) and len(v.strip())>3:
                t=_norm(v).split(' - ')[0].strip()
                if t and not any(x in t for x in ('AGOSTO','JULHO','SETEMBRO','COMPETENCIA')):
                    title=_canonical(t); break
        if title: break
    # Se o título tem alias conhecido, prefira-o; caso contrário preserve a aba.
    return title if title in set(ALIASES.values()) else sheet_key

def _excel_date(v):
    if isinstance(v,(datetime,date)):
        return v.strftime('%d/%m/%Y')
    if isinstance(v,(int,float)) and 40000 <= float(v) <= 60000:
        try: return (datetime(1899,12,30)+timedelta(days=float(v))).strftime('%d/%m/%Y')
        except Exception: return None
    return None

def read_sadt(path):
    wb=load_workbook(path, data_only=True, read_only=True)
    result={}
    for s in wb.sheetnames:
        if _norm(s) in IGNORE_SHEETS or _norm(s).startswith('GRAF'):
            continue
        ws=wb[s]
        exam=_exam_name(s,ws)
        rows=list(ws.iter_rows(values_only=True))
        daily=defaultdict(lambda:{'interno':0,'externo':0,'total':0})
        found_block=False

        # Cada planilha tem dois blocos (1-15 e 16-31). Lê ambos, sem depender da coluna "Total".
        for r_idx,row in enumerate(rows):
            date_cols={}
            for c,v in enumerate(row):
                d=_excel_date(v)
                if d: date_cols[c]=d
            if len(date_cols)<2:
                continue
            for rr in range(r_idx+1,min(r_idx+5,len(rows))):
                label=_norm(' '.join(str(x or '') for x in rows[rr][:4]))
                mode='interno' if 'INTERNO' in label else ('externo' if 'EXTERNO' in label else None)
                if not mode: continue
                found_block=True
                for c,d in date_cols.items():
                    if c>=len(rows[rr]): continue
                    v=rows[rr][c]
                    if isinstance(v,(int,float)) and not isinstance(v,bool):
                        q=max(0,int(v))
                        daily[d][mode]+=q
                        daily[d]['total']+=q

        # Total calculado pelos dias evita depender de fórmulas/mesclas inconsistentes da planilha.
        ti=sum(v['interno'] for v in daily.values())
        te=sum(v['externo'] for v in daily.values())
        if found_block:
            result[exam]={
                'sheet':s,'interno':ti,'externo':te,'total':ti+te,
                'daily':dict(daily),'found':True,
            }
    return result

def slice_period(data, start_ddmmyyyy, end_ddmmyyyy):
    if not start_ddmmyyyy or not end_ddmmyyyy: return data
    try:
        a=datetime.strptime(start_ddmmyyyy,'%d/%m/%Y').date(); b=datetime.strptime(end_ddmmyyyy,'%d/%m/%Y').date()
    except Exception: return data
    out={}
    for exam,r in data.items():
        daily={}; ti=te=0
        for ds,v in r.get('daily',{}).items():
            try:d=datetime.strptime(ds,'%d/%m/%Y').date()
            except:continue
            if a<=d<=b:
                daily[ds]=v; ti+=int(v.get('interno',0)); te+=int(v.get('externo',0))
        out[exam]={**r,'daily':daily,'interno':ti,'externo':te,'total':ti+te}
    return out
